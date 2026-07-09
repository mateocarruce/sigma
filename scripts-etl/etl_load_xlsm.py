"""
ETL: MATRIZ_FACTURACION_FINAL8_0.xlsm  ->  PostgreSQL

Carga tres fuentes reales encontradas en el archivo del hospital:
  1. Hoja 'MATRIZ'            -> tabla planillas (+ pacientes anonimizados, lotes_carga)
  2. Hoja 'AYUDANTIA > 2023'  -> tabla tarifario_nacional (bloque de códigos/valores)
  3. Hoja 'AYUDANTIA > 2023'  -> tabla catalogo_motivos_objecion (bloque de glosas CMT/LQD/REV)

Uso:
    python etl_load_xlsm.py --file /ruta/MATRIZ_FACTURACION_FINAL8_0.xlsm

Variables de entorno esperadas (coinciden con docker-compose.yml):
    DB_HOST (default: localhost)
    DB_PORT (default: 5432)
    DB_NAME (default: glosas_db)
    DB_USER (default: hospital_admin)
    DB_PASSWORD
    HASH_SALT  -> salt para anonimizar cédulas (NUNCA hardcodear en producción)
"""

import argparse
import hashlib
import os
import sys
from datetime import datetime

import openpyxl
import psycopg2
import psycopg2.extras
from psycopg2.extras import execute_values

# ------------------------------------------------------------------
# Configuración
# ------------------------------------------------------------------

DB_CONFIG = {
    "host": os.getenv("DB_HOST", "localhost"),
    "port": os.getenv("DB_PORT", "5432"),
    "dbname": os.getenv("DB_NAME", "glosas_db"),
    "user": os.getenv("DB_USER", "hospital_admin"),
    "password": os.getenv("DB_PASSWORD", "changeme_en_env"),
}

HASH_SALT = os.getenv("HASH_SALT", "cambiar_este_salt_en_produccion")


def hash_identificador(cedula: str) -> str:
    """Anonimiza la cédula/pasaporte con SHA-256 + salt (cumplimiento LOPDP)."""
    valor = f"{HASH_SALT}:{str(cedula).strip()}"
    return hashlib.sha256(valor.encode("utf-8")).hexdigest()


def rango_edad_placeholder() -> str:
    # El archivo actual no trae fecha de nacimiento, solo cédula y nombre.
    # Se deja el campo listo para cuando se incorpore esa fuente.
    return None


# ------------------------------------------------------------------
# 1. Parseo del Tarifario Nacional (hoja AYUDANTIA > 2023, bloque superior)
# ------------------------------------------------------------------

def parse_tarifario(ws):
    """
    Extrae filas del tarifario. Se detienen automáticamente al llegar al
    bloque de 'MOTIVOS DE OBJECIÓN...' (columna B), que marca el inicio
    del catálogo de glosas más abajo en la misma hoja.
    """
    registros = []
    en_bloque_objeciones = False

    for row in ws.iter_rows(min_row=1, values_only=True):
        a, b, c, d, e = row[0], row[1], row[2], row[3], row[4]

        if isinstance(b, str) and b.strip().upper().startswith("MOTIVOS DE OBJECI"):
            en_bloque_objeciones = True
            continue

        if en_bloque_objeciones:
            continue

        # Fila válida de tarifario: especialidad (texto), código (numérico), descripción
        if isinstance(a, str) and isinstance(c, (int, float)) and d:
            registros.append(
                (
                    a.strip(),
                    str(b).strip() if b is not None else None,
                    str(c).strip(),
                    str(d).strip(),
                    float(e) if isinstance(e, (int, float)) else 0.0,
                )
            )

    return registros


# ------------------------------------------------------------------
# 2. Parseo del Catálogo de Motivos de Objeción (mismo sheet, bloque inferior)
# ------------------------------------------------------------------

CATEGORIA_MAP = {
    "MOTIVOS DE OBJECIÓNOBJECIONES DE PERTINENCIA MÉDICA": "PERTINENCIA_MEDICA",
    "MOTIVOS DE OBJECIÓN CONTROL DE TARIFAS": "CONTROL_TARIFAS",
    "MOTIVOS DE OBJECIÓN EN REVISIÓN DOCUMENTAL": "REVISION_DOCUMENTAL",
}


def parse_motivos_objecion(ws):
    registros = []
    categoria_actual = None
    vistos = set()

    for row in ws.iter_rows(min_row=1, values_only=True):
        a, b = row[0], row[1]

        if isinstance(b, str) and b.strip().upper().startswith("MOTIVOS DE OBJECI"):
            texto = b.strip()
            categoria_actual = CATEGORIA_MAP.get(texto)  # None para el bloque "CHECK LIST" vacío
            continue

        if categoria_actual is None:
            continue

        if b and (isinstance(a, str) or a == 0):
            codigo = "0" if a == 0 else str(a).strip()
            descripcion = str(b).strip()
            # 'codigo' es UNIQUE a nivel de tabla (no por categoría), así que
            # si el mismo código aparece en más de una sección del Excel
            # (ej. '0' = SIN OBJECIÓN repetido), solo se conserva la primera vez.
            if codigo in vistos:
                continue
            vistos.add(codigo)
            registros.append((codigo, descripcion, categoria_actual))

    return registros


# ------------------------------------------------------------------
# 3. Parseo de la MATRIZ (planillas reales)
# ------------------------------------------------------------------

def parse_matriz(ws):
    """
    Header está en la fila 8. Los datos empiezan en la fila 9.
    Columnas (1-indexed): A=#, B=Nombre, C=Cod.Validación, D=CI, E=CIE10,
    F=Fecha, G=Especialidad/Honorario, H=Codigo TPSNS, I=Nivel, J=Descripción,
    K=Desc.Medicamento, L=Cantidad, M=Valor Unitario TPSNS, N=Valor Unit.Medic,
    O=Subtotal, ..., W=Valor Solicitado, X=Servicio, Y=Sin Objeción (ground truth futuro)
    """
    registros = []
    for row in ws.iter_rows(min_row=9, max_col=25, values_only=True):
        if not any(v not in (None, "") for v in row):
            continue

        (
            _num, nombre, _cod_val, cedula, cie10, fecha, especialidad,
            cod_tpsns, nivel, descripcion, _desc_med, cantidad,
            _val_unit_tpsns, _val_unit_med, _subtotal, *resto
        ) = row

        valor_solicitado = row[22] if len(row) > 22 else None
        servicio = row[23] if len(row) > 23 else None
        motivo_real = row[24] if len(row) > 24 else None

        if cedula is None or servicio is None:
            # Filas de encabezado/firma que no son registros de facturación real
            continue

        registros.append(
            {
                "hash_identificador": hash_identificador(cedula),
                "tipo_seguro": None,  # no viene explícito en esta hoja; se infiere en NestJS si aplica
                "numero_factura": str(cedula),  # placeholder: el archivo no trae nro. de factura formal
                "codigo_cie10": str(cie10).strip() if cie10 else None,
                "codigo_prestacion": str(cod_tpsns).strip() if cod_tpsns else None,
                "especialidad": str(especialidad).strip() if especialidad else None,
                "descripcion": str(descripcion).strip() if descripcion else None,
                "cantidad": float(cantidad) if isinstance(cantidad, (int, float)) else None,
                "valor_facturado": float(valor_solicitado) if isinstance(valor_solicitado, (int, float)) else 0.0,
                "fecha_atencion": fecha if isinstance(fecha, datetime) else None,
                "tipo_servicio": str(servicio).strip() if servicio else None,
                "motivo_objecion_real": str(motivo_real).strip() if motivo_real else None,
            }
        )
    return registros


# ------------------------------------------------------------------
# 4. Carga a PostgreSQL
# ------------------------------------------------------------------

def cargar_tarifario(conn, registros):
    sql = """
        INSERT INTO tarifario_nacional
            (especialidad, nivel, codigo_prestacion, descripcion, valor_unitario_oficial)
        VALUES %s
        ON CONFLICT (especialidad, nivel, codigo_prestacion) DO NOTHING
    """
    with conn.cursor() as cur:
        execute_values(cur, sql, registros, page_size=5000)
    conn.commit()
    print(f"[tarifario_nacional] {len(registros)} filas procesadas.")


def cargar_motivos_objecion(conn, registros):
    sql = """
        INSERT INTO catalogo_motivos_objecion (codigo, descripcion, categoria)
        VALUES %s
        ON CONFLICT (codigo) DO UPDATE SET
            descripcion = EXCLUDED.descripcion,
            categoria = EXCLUDED.categoria
    """
    with conn.cursor() as cur:
        execute_values(cur, sql, registros, page_size=500)
    conn.commit()
    print(f"[catalogo_motivos_objecion] {len(registros)} filas procesadas.")


def cargar_planillas(conn, registros, nombre_archivo):
    with conn.cursor() as cur:
        # 1. Crear el lote de carga (representa "un archivo subido por el auditor")
        cur.execute(
            """
            INSERT INTO lotes_carga (nombre_archivo, total_registros, estado)
            VALUES (%s, %s, 'PROCESANDO')
            RETURNING id
            """,
            (nombre_archivo, len(registros)),
        )
        lote_id = cur.fetchone()[0]

        errores = 0
        for r in registros:
            # 2. Upsert de paciente anonimizado
            cur.execute(
                """
                INSERT INTO pacientes (hash_identificador, tipo_seguro)
                VALUES (%s, %s)
                ON CONFLICT (hash_identificador) DO NOTHING
                RETURNING id
                """,
                (r["hash_identificador"], r["tipo_seguro"]),
            )
            paciente_row = cur.fetchone()
            if paciente_row:
                paciente_id = paciente_row[0]
            else:
                cur.execute(
                    "SELECT id FROM pacientes WHERE hash_identificador = %s",
                    (r["hash_identificador"],),
                )
                paciente_id = cur.fetchone()[0]

            # 3. Validación mínima de campos críticos (reglas de negocio simples)
            reglas_detalle = []
            if not r["codigo_cie10"]:
                reglas_detalle.append("CIE-10 ausente")
            if not r["valor_facturado"] or r["valor_facturado"] == 0:
                reglas_detalle.append("Valor facturado en cero")
            if not r["codigo_prestacion"]:
                reglas_detalle.append("Código de prestación ausente")

            reglas_estado = "ERROR_CRITICO" if reglas_detalle else "OK"
            if reglas_detalle:
                errores += 1

            # 4. Insertar planilla
            cur.execute(
                """
                INSERT INTO planillas (
                    lote_id, paciente_id, numero_factura, codigo_cie10,
                    codigo_prestacion, especialidad, valor_facturado, cantidad,
                    fecha_atencion, tipo_seguro, reglas_estado, reglas_detalle,
                    estado_planilla, motivo_objecion_real
                ) VALUES (
                    %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s
                )
                """,
                (
                    lote_id, paciente_id, r["numero_factura"], r["codigo_cie10"],
                    r["codigo_prestacion"], r["especialidad"], r["valor_facturado"], r["cantidad"],
                    r["fecha_atencion"], r["tipo_servicio"], reglas_estado,
                    psycopg2.extras.Json(reglas_detalle), "CARGADA", r["motivo_objecion_real"],
                ),
            )

        cur.execute(
            """
            UPDATE lotes_carga
            SET registros_con_error = %s, estado = 'VALIDADO'
            WHERE id = %s
            """,
            (errores, lote_id),
        )

    conn.commit()
    print(f"[planillas] {len(registros)} filas cargadas en lote_id={lote_id} ({errores} con error de reglas).")


# ------------------------------------------------------------------
# Main
# ------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(description="ETL .xlsm -> PostgreSQL")
    parser.add_argument("--file", required=True, help="Ruta al archivo .xlsm")
    parser.add_argument("--skip-tarifario", action="store_true", help="Omitir carga del tarifario (rápido para pruebas)")
    args = parser.parse_args()

    print(f"Leyendo workbook: {args.file}")
    wb = openpyxl.load_workbook(args.file, data_only=True, read_only=True)

    conn = psycopg2.connect(**DB_CONFIG)

    try:
        ws_ayudantia = wb["AYUDANTIA > 2023"]

        if not args.skip_tarifario:
            print("Parseando tarifario nacional (esto puede tardar unos segundos, ~89k filas)...")
            tarifario = parse_tarifario(ws_ayudantia)
            cargar_tarifario(conn, tarifario)
        else:
            print("Tarifario omitido (--skip-tarifario).")

        print("Parseando catálogo de motivos de objeción...")
        motivos = parse_motivos_objecion(ws_ayudantia)
        cargar_motivos_objecion(conn, motivos)

        print("Parseando MATRIZ (planillas)...")
        ws_matriz = wb["MATRIZ"]
        planillas = parse_matriz(ws_matriz)
        if planillas:
            cargar_planillas(conn, planillas, os.path.basename(args.file))
        else:
            print("[planillas] No se encontraron filas de datos en MATRIZ.")

        print("ETL completado con éxito.")

    except Exception as exc:
        conn.rollback()
        print(f"ERROR durante el ETL, se hizo rollback: {exc}", file=sys.stderr)
        raise
    finally:
        conn.close()


if __name__ == "__main__":
    main()
